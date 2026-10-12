"""`--export-excel [FILE]`: the dashboard's jobs (per track), your applications, and coverage as one .xlsx."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from ..output.dashboard import select_visible
from ..storage import db
from ..tracker import statuses as S
from ..tracker.applications import STAGE_LABELS

JOB_COLS = [("Company", 22), ("Title", 48), ("Label", 9), ("Score", 7), ("JD match", 9), ("ATS %", 7), ("Level", 9),
            ("Location", 30), ("Work mode", 10), ("Posted", 11), ("Deadline", 11), ("Sponsorship", 16),
            ("My status", 14), ("Source", 16), ("Apply link", 40)]


def _safe(v):
    """Text never becomes a formula in Excel."""
    return ("'" + v) if isinstance(v, str) and v[:1] in ("=", "+", "-", "@") else v


def _sheet(wb, title, cols, rows, link_col=None):
    ws = wb.create_sheet(title)
    ws.append([c for c, _ in cols])
    for c in ws[1]:
        c.font = Font(bold=True)
    for i, (_, w) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in rows:
        ws.append([_safe(v) for v in r])
        if link_col is not None and isinstance(r[link_col], str) and r[link_col].startswith("http"):
            cell = ws.cell(row=ws.max_row, column=link_col + 1)
            cell.hyperlink, cell.style = r[link_col], "Hyperlink"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def export(conn, path: Path, top_n: int) -> Path:
    jobs = select_visible(db.open_jobs(conn), top_n)
    wb = Workbook()
    wb.remove(wb.active)
    for track, title in (("coop_intern", "Co-op & Internships"), ("full_time", "Full-time")):
        rows = []
        for j in sorted(jobs, key=lambda j: -(j.get("score") or 0)):
            t = j.get("track_override") or j.get("track") or "full_time"
            if t not in (track, "unclear"):
                continue
            rows.append([j["company"], j["title"], j.get("label"), round(j.get("score") or 0), round(j.get("qual_score") or 0),
                         round(j["ats_coverage"]) if j.get("ats_coverage") is not None else None, j.get("level"),
                         j.get("location"), j.get("work_mode"), (j.get("posted_date") or "")[:10] or "Unknown",
                         j.get("deadline"), ((j.get("flags") or {}).get("sponsorship") or {}).get("value", "Not stated"),
                         j.get("user_status") or "", j["source"], j.get("apply_url") or j.get("source_url")])
        _sheet(wb, title, JOB_COLS, rows, link_col=14)
    apps = S.applications(conn, include_deleted=False, with_jd=False)
    _sheet(wb, "Applications", [("Company", 22), ("Role", 44), ("Track", 12), ("Stage", 18), ("Applied", 11),
                                ("Next step", 30), ("Next step date", 13), ("Contact", 20), ("Referral", 18),
                                ("Source", 16), ("Link", 40), ("Notes", 50)],
           [[a["company"], a["title"], a["track"], STAGE_LABELS.get(a["stage"], a["stage"]), a["applied_date"],
             a["next_step"], a["next_step_date"], a["contact"], a["referral"], a["source"], a["link"], a["notes"]]
            for a in apps], link_col=10)
    _sheet(wb, "Coverage", [("Company", 28), ("Source", 18), ("Status", 10), ("Last checked", 22), ("Jobs listed", 11),
                            ("Notes", 70)],
           [[r["company"], r["source"], r["status"], r["last_checked"], r["job_count"], r["error"]]
            for r in db.source_statuses(conn)])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
