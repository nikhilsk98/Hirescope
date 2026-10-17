"""My Applications as an Excel file on the Desktop, regenerated from SQLite after every change:
~/Desktop/Generated Resumes/HireScope Applications.xlsx (preferences.yaml → excel_applications.path / enabled).

Sheets: Applications (one row per application, newest first, clickable job / resume / saved-JD links, stage colors),
Applied Companies (one row per company), Summary (totals by stage and track, applications per week, response rate).

Written after a dashboard change reaches SQLite through the launcher (Mark Applied, stage changes, notes, the Bucket),
after every HireScope run, import_pending.sh, `resume_builder refresh` and every queue build.

Safe writes: the workbook is built in data.nosync, then a child process (ops/desktop_access, killed after a timeout so
a macOS permission prompt can never hang a background job) copies it next to the old one and replaces it atomically. If the
file is open in Excel (its "~$" owner file is there, or Excel holds it open), the old file is left alone and the update is
saved as "HireScope Applications (pending).xlsx"; the next sync (once Excel has closed it) replaces the file and
removes the pending copy. The dashboard shows a notice meanwhile. Nothing is ever deleted from SQLite here.
"""
from __future__ import annotations

import json
import os
import subprocess
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .. import paths
from ..storage import db
from ..tracker import applications as A
from ..tracker import statuses as S

STAGE_FILL = {"ready": "FFF4CC", "applied": "DCE8FB", "oa": "E6E0FA", "interviewing": "D9F2E6", "offer": "B9EBC9",
              "rejected": "F6D6D3", "ghosted": "E7E7E7", "withdrawn": "EFEFEF"}
APP_COLS = [("Date applied", 12), ("Company", 24), ("Role", 44), ("Track", 13), ("Stage", 18), ("Source", 11),
            ("Location", 26), ("Job ID", 9), ("Job link", 14), ("Resume PDF", 16), ("Resume version", 9),
            ("Saved JD", 11), ("Follow-up due", 13), ("Days since applied", 9), ("Notes", 50), ("Last updated", 18)]
TRACKS = {"coop_intern": "Co-op / Intern", "full_time": "Full-time"}


def settings() -> tuple[bool, Path]:
    env = os.environ.get("JOBSEARCH_EXCEL_PATH")
    try:
        from ..config import load_preferences
        cfg = load_preferences().excel_applications
        enabled, path = cfg.enabled, cfg.path
    except (OSError, ValueError):
        enabled, path = True, "~/Desktop/Generated Resumes/HireScope Applications.xlsx"
    return enabled, Path(os.path.expanduser(env or path))


def pending_path(path: Path) -> Path:
    return path.with_name(f"{path.stem} (pending){path.suffix}")


def _safe(v):
    """Text never becomes a formula in Excel."""
    return ("'" + v) if isinstance(v, str) and v[:1] in ("=", "+", "-", "@") else v


def _link(cell, target: str | None, text: str) -> None:
    if not target:
        return
    cell.value = text
    cell.hyperlink = target
    cell.style = "Hyperlink"


def _file_url(p: str | Path | None) -> str | None:
    if not p:
        return None
    p = Path(p)
    return p.as_uri() if p.is_absolute() and p.exists() else None


def desktop_pdf(conn, app: dict) -> str | None:
    """The Desktop copy of the resume version the application recorded (current → <stem>.pdf, older → _v<N>.pdf);
    the data.nosync PDF when there is no Desktop copy."""
    ref = (app.get("job") or {}).get("resume") or {}
    pk, ver, pdf = ref.get("job_pk"), ref.get("version"), ref.get("pdf")
    if not pk or not ver:
        return None
    try:
        from resume_builder import desktop, pipeline, versions
        j = conn.execute("SELECT company, title FROM jobs WHERE id=?", (pk,)).fetchone()
        if j:
            jd = pipeline.job_dir(j["company"], pk, j["title"])
            enabled, folder = desktop.settings_of()
            cur = (versions.current(jd) or {}).get("version")
            name = Path(pdf or "").name
            if enabled and name:
                cand = desktop.target_dir(jd, folder) / (name if cur == ver else f"{Path(name).stem}_v{ver}.pdf")
                if cand.exists():
                    return cand.as_uri()
    except Exception:  # noqa: BLE001 — a link is a convenience; the row is still written
        pass
    return _file_url(pdf)


def build(conn, today: date | None = None) -> Workbook:
    from ..config import load_preferences
    try:
        window = tuple(load_preferences().applications.follow_up_days)
    except (OSError, ValueError):
        window = (7, 10)
    today = today or date.today()
    apps = S.applications(conn, include_deleted=False, with_jd=False)
    jobs = {r["id"]: r for r in conn.execute("SELECT id, location, job_id, apply_url, source_url FROM jobs WHERE id IN "
                                             "(SELECT job_pk FROM applications WHERE job_pk IS NOT NULL)")}
    jd_paths = {r["id"]: r["jd_path"] for r in conn.execute("SELECT id, jd_path FROM applications")}
    apps.sort(key=lambda a: (a.get("applied_date") or (a.get("updated_at") or "")[:10], a.get("updated_at") or ""),
              reverse=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Applications"
    ws.append([c for c, _ in APP_COLS])
    for a in apps:
        j = jobs.get(a.get("job_pk")) if a.get("job_pk") else None
        if j is None and a.get("job_key"):
            pk = S.find_job(conn, a["job_key"])
            j = jobs.get(pk) or (conn.execute("SELECT id, location, job_id, apply_url, source_url FROM jobs WHERE id=?",
                                              (pk,)).fetchone() if pk else None)
        fu = A.follow_up(a, today, window)
        d = A._d(a.get("applied_date"))
        ref = (a.get("job") or {}).get("resume") or {}
        ws.append([_safe(x) for x in (
            a.get("applied_date") or "", a.get("company"), a.get("title"), TRACKS.get(a.get("track"), a.get("track") or ""),
            A.STAGE_LABELS.get(a.get("stage"), a.get("stage")), a.get("source") or "", (j["location"] if j else "") or "",
            j["id"] if j else "", "", "", ref.get("version") or "", "", fu["due"] if fu else "",
            (today - d).days if d else "", a.get("notes") or "", (a.get("updated_at") or "")[:19].replace("T", " "))])
        r = ws.max_row
        link = a.get("link") or (j and (j["apply_url"] or j["source_url"])) or ""
        _link(ws.cell(r, 9), link if str(link).startswith("http") else None, "Open job")
        _link(ws.cell(r, 10), desktop_pdf(conn, a), "Resume PDF")
        _link(ws.cell(r, 12), _file_url(paths.stored_file(jd_paths.get(a["id"]), paths.APPS_DIR)), "Saved JD")
        fill = STAGE_FILL.get(a.get("stage"))
        if fill:
            ws.cell(r, 5).fill = PatternFill("solid", fgColor=fill)
        ws.cell(r, 15).alignment = Alignment(wrap_text=True, vertical="top")
    _finish(ws, APP_COLS)

    wc = wb.create_sheet("Applied Companies")
    cols = [("Company", 26), ("Applications", 13), ("Latest stage", 18), ("Latest date", 12), ("Sources", 18),
            ("Latest link", 14), ("All applications", 70)]
    wc.append([c for c, _ in cols])
    for row in A.applied_companies(apps):
        latest = row["applications"][0] if row["applications"] else {}
        wc.append([_safe(row["company"]), row["count"], A.STAGE_LABELS.get(row["latest_stage"], row["latest_stage"]),
                   row["latest_applied"], ", ".join(row["sources"]), "",
                   _safe("; ".join(f"{x['title']} ({A.STAGE_LABELS.get(x['stage'], x['stage'])}, {x['applied_date'] or '?'})"
                                   + (f" {x['link']}" if x["link"] else "") for x in row["applications"]))])
        r = wc.max_row
        _link(wc.cell(r, 6), latest.get("link") if str(latest.get("link") or "").startswith("http") else None, "Open job")
        if (fill := STAGE_FILL.get(row["latest_stage"])):
            wc.cell(r, 3).fill = PatternFill("solid", fgColor=fill)
        wc.cell(r, 7).alignment = Alignment(wrap_text=True, vertical="top")
    _finish(wc, cols)

    sm = wb.create_sheet("Summary")
    sm.column_dimensions["A"].width = 30
    for c in "BCDE":
        sm.column_dimensions[c].width = 16
    bold = Font(bold=True)
    sm.append(["HireScope applications", f"updated {datetime.now().strftime('%Y-%m-%d %H:%M')}"])
    sm["A1"].font = Font(bold=True, size=13)
    sm.append([])
    sm.append(["By stage", "Applications"])
    sm.cell(sm.max_row, 1).font = sm.cell(sm.max_row, 2).font = bold
    for st in A.STAGES:
        n = sum(1 for a in apps if a.get("stage") == st)
        sm.append([A.STAGE_LABELS[st], n])
        if (fill := STAGE_FILL.get(st)):
            sm.cell(sm.max_row, 1).fill = PatternFill("solid", fgColor=fill)
    real = [a for a in apps if A.is_applied(a)]
    sm.append(["Total applied (excl. Ready to apply)", len(real)])
    sm.cell(sm.max_row, 1).font = bold
    sm.append([])
    sm.append(["By track", "Applications", "Response rate", "Interview rate", "Offers"])
    for c in range(1, 6):
        sm.cell(sm.max_row, c).font = bold
    for t, label in TRACKS.items():
        st = A.track_stats([a for a in apps if a.get("track") == t], today)
        sm.append([label, st["total"], f"{st['response_rate']}%" if st["response_rate"] is not None else "—",
                   f"{st['interview_rate']}%" if st["interview_rate"] is not None else "—", st["offers"]])
    allst = A.track_stats(apps, today)
    sm.append(["All", allst["total"], f"{allst['response_rate']}%" if allst["response_rate"] is not None else "—",
               f"{allst['interview_rate']}%" if allst["interview_rate"] is not None else "—", allst["offers"]])
    sm.cell(sm.max_row, 1).font = bold
    sm.append([])
    sm.append(["Applications per week (week of)", "Applications"])
    sm.cell(sm.max_row, 1).font = sm.cell(sm.max_row, 2).font = bold
    start = date.fromisoformat(allst["week_start"])
    for i, n in enumerate(allst["per_week"]):
        sm.append([(start + timedelta(weeks=i)).isoformat(), n])
    sm.append([])
    sm.append(["A response is any move to Online Assessment, Interviewing, Offer or Rejected."])
    return wb


def _finish(ws, cols) -> None:
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="EDEFF3")
    for i, (_, w) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(ws.max_row, 1)}"


def open_in_excel(path: Path) -> bool:
    """Excel has the file open: its "~$" owner file sits next to it, or an Excel process holds it."""
    owner = [path.with_name("~$" + path.name), path.with_name("~$" + path.name[2:])]
    if any(o.exists() for o in owner):
        return True
    if os.environ.get("HIRESCOPE_NO_LSOF") or not path.exists():
        return False
    try:
        out = subprocess.run(["lsof", "-Fc", "--", str(path)], capture_output=True, text=True, timeout=5).stdout
    except (OSError, subprocess.SubprocessError):
        return False
    return "Microsoft Excel" in out or "cExcel" in out


def _record(conn, **state) -> None:
    db.set_meta(conn, "excel_applications", json.dumps({**state, "updated_at": datetime.now(timezone.utc)
                                                        .isoformat(timespec="seconds")}))
    conn.commit()


def update(conn, path: Path | None = None) -> str:
    """Regenerate the workbook (atomic replace, or the pending copy while the file is open). Returns a summary line."""
    enabled, default = settings()
    path = path or default
    if not enabled:
        _record(conn, status="off", path=str(path))
        return ""
    wb = build(conn)
    n = wb["Applications"].max_row - 1
    staged = paths.STATE_DIR / "excel" / path.name          # built in data.nosync; only a child touches the Desktop
    staged.parent.mkdir(parents=True, exist_ok=True)
    wb.save(staged)
    from .desktop_access import place
    res = place(staged, path)
    st = res.get("status", "error")
    if st == "ok":
        _record(conn, status="ok", path=str(path), message=f"{n} application(s)")
        return ""
    _record(conn, status=st, path=str(path), pending_path=res.get("pending_path"), message=res.get("message", ""))
    return res.get("message", "")
