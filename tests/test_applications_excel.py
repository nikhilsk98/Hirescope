"""HireScope Applications.xlsx on the Desktop: rows, links, three sheets, stage colors, atomic writes, the
file-open fallback, instant sync through the launcher. DEMO data in a temporary folder."""
import json
import os
from pathlib import Path

import pytest
from openpyxl import load_workbook

from jobsearch import launcher, paths
from jobsearch.ops import applications_excel as X
from jobsearch.storage import db

from .resume_builder.conftest import insert_job


@pytest.fixture
def xconn():
    c = db.connect(paths.DB)
    a = insert_job(c, paths.DATA / "jd", company="DemoCorp", title="Backend Engineer Intern", job_id="x-1")
    b = insert_job(c, paths.DATA / "jd", company="=HYPERLINK(evil)", title="Software Engineer I", track="full_time",
                   job_id="x-2")
    c.execute("INSERT INTO resume_builds (job_pk, version, status, pdf_path, manifest_path, created_at, summary) VALUES "
              "(?,2,'ready','/nonexistent/v2/R.pdf','/nonexistent/v2/manifest.json','2026-10-16T00:00:00+00:00','{}')", (a,))
    c.execute("INSERT INTO resume_status (job_pk, status, current_version, latest_version, updated_at) VALUES "
              "(?, 'ready', 2, 2, '2026-10-16T00:00:00+00:00')", (a,))
    c.commit()
    yield c, a, b
    c.close()


def xlsx() -> Path:
    return X.settings()[1]


def rows(sheet="Applications"):
    ws = load_workbook(xlsx())[sheet]
    return [[c.value for c in r] for r in ws.iter_rows()], ws


def test_applications_sheet_rows_links_and_order(xconn):
    conn, a, b = xconn
    launcher.handle(f"hirescope://status?job={a}&stage=applied&app=aA&date=2026-10-15&v=2")
    launcher.handle(f"hirescope://status?job={b}&stage=applied&app=aB&date=2026-10-17")
    data, ws = rows()
    assert data[0] == [c for c, _ in X.APP_COLS]
    assert [r[1] for r in data[1:]] == ["'=HYPERLINK(evil)", "DemoCorp"]           # newest first; formula-safe text
    demo = next(r for r in data if r[1] == "DemoCorp")
    assert demo[0] == "2026-10-15" and demo[3] == "Co-op / Intern" and demo[4] == "Applied" and demo[5] == "HireScope"
    assert demo[6] == "Boston, MA" and demo[7] == a and demo[10] == 2 and demo[12] == "2026-10-22"
    r = next(i for i, row in enumerate(data, 1) if row[1] == "DemoCorp")
    assert ws.cell(r, 9).hyperlink.target == "https://example.com/apply" and ws.cell(r, 9).value == "Open job"
    assert ws.cell(r, 5).fill.fgColor.rgb.endswith("DCE8FB")                       # stage color
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref.startswith("A1:") and ws["A1"].font.bold
    assert ws.column_dimensions["C"].width >= 40
    assert not list(xlsx().parent.glob(".*.tmp.xlsx"))                               # atomic: no temp file left


def test_saved_jd_and_desktop_resume_links(xconn, tmp_path, monkeypatch):
    conn, a, b = xconn
    launcher.handle(f"hirescope://status?job={a}&stage=applied&app=aA&v=2")
    data, ws = rows()
    r = next(i for i, row in enumerate(data, 1) if row[1] == "DemoCorp")
    assert ws.cell(r, 12).hyperlink.target.startswith("file://") and ws.cell(r, 12).value == "Saved JD"
    # the Desktop copy of that exact version, when it exists
    from resume_builder import desktop, pipeline
    jd = pipeline.job_dir("DemoCorp", a, "Backend Engineer Intern")
    jd.mkdir(parents=True, exist_ok=True)
    (jd / "current.json").write_text(json.dumps({"version": 2, "status": "ready"}))
    folder = desktop.settings_of()[1]
    dst = desktop.target_dir(jd, folder) / "R.pdf"
    dst.parent.mkdir(parents=True, exist_ok=True)
    dst.write_bytes(b"%PDF DEMO")
    X.update(conn)
    data, ws = rows()
    assert ws.cell(r, 10).hyperlink.target == dst.as_uri()


def test_companies_and_summary_sheets(xconn):
    conn, a, b = xconn
    launcher.handle(f"hirescope://status?job={a}&stage=applied&app=aA&date=2026-10-15")
    launcher.handle("hirescope://status?app=aA&stage=interviewing")
    launcher.handle(f"hirescope://status?job={b}&stage=applied&app=aB&date=2026-10-17")
    wb = load_workbook(xlsx())
    assert wb.sheetnames == ["Applications", "Applied Companies", "Summary"]
    comp = [[c.value for c in r] for r in wb["Applied Companies"].iter_rows()]
    assert comp[0][:4] == ["Company", "Applications", "Latest stage", "Latest date"]
    assert {r[0] for r in comp[1:]} == {"DemoCorp", "'=HYPERLINK(evil)"}
    assert next(r for r in comp if r[0] == "DemoCorp")[2] == "Interviewing"
    summ = [[c.value for c in r] for r in wb["Summary"].iter_rows()]
    flat = {r[0]: r[1:] for r in summ if r and r[0]}
    assert flat["Interviewing"][0] == 1 and flat["Applied"][0] == 1 and flat["Total applied (excl. Ready to apply)"][0] == 2
    assert flat["Co-op / Intern"][0] == 1 and flat["Full-time"][0] == 1 and flat["All"][1] == "50.0%"
    assert any(str(r[0]).startswith("Applications per week") for r in summ if r)


def test_stage_change_and_notes_sync_instantly(xconn):
    conn, a, b = xconn
    launcher.handle(f"hirescope://status?job={a}&stage=applied&app=aA")
    from urllib.parse import quote
    launcher.handle("hirescope://note?app=aA&text=" + quote("Recruiter: Dana\nfollow up Monday"))
    launcher.handle("hirescope://status?app=aA&stage=oa")
    data, ws = rows()
    demo = next(r for r in data if r[1] == "DemoCorp")
    assert demo[4] == "Online Assessment" and demo[14] == "Recruiter: Dana\nfollow up Monday"
    launcher.handle(f"hirescope://note?job={a}&text=" + quote("liked the team"))
    assert conn.execute("SELECT user_notes FROM jobs WHERE id=?", (a,)).fetchone()[0] == "liked the team"
    for bad in ("hirescope://note?text=x", f"hirescope://note?job={a}&app=aA&text=x", "hirescope://note?app=aA"):
        with pytest.raises(launcher.LinkError):
            launcher.parse(bad)


def test_open_in_excel_saves_pending_and_applies_next_sync(xconn):
    conn, a, b = xconn
    launcher.handle(f"hirescope://status?job={a}&stage=applied&app=aA")
    before = xlsx().read_bytes()
    owner = xlsx().with_name("~$" + xlsx().name)
    owner.write_text("Excel owner file (DEMO)")
    launcher.handle("hirescope://status?app=aA&stage=interviewing")
    assert xlsx().read_bytes() == before                                    # the open file is never touched
    pend = X.pending_path(xlsx())
    assert pend.exists() and load_workbook(pend)["Applications"]["E2"].value == "Interviewing"
    st = json.loads(db.get_meta(conn, "excel_applications"))
    assert st["status"] == "pending" and "open in Excel" in st["message"]
    live = (paths.DASHBOARD / "live.js").read_text()
    assert '"status": "pending"' in live                                   # the dashboard shows a notice
    owner.unlink()                                                          # Excel closed
    launcher.handle("hirescope://status?app=aA&stage=offer")
    assert not pend.exists() and load_workbook(xlsx())["Applications"]["E2"].value == "Offer"
    assert json.loads(db.get_meta(conn, "excel_applications"))["status"] == "ok"


def test_replace_failure_keeps_the_old_file(xconn, monkeypatch, capsys):
    """The child that places the file: a failed replace keeps the old file and saves the update beside it."""
    from jobsearch.ops import desktop_access
    conn, a, b = xconn
    X.update(conn)
    before = xlsx().read_bytes()
    staged = paths.STATE_DIR / "excel" / xlsx().name
    real = Path.replace

    def boom(self, target):
        if Path(target) == xlsx():
            raise PermissionError(13, "Permission denied")
        return real(self, target)
    monkeypatch.setattr(Path, "replace", boom)
    assert desktop_access._place_main(str(staged), str(xlsx())) == 0
    out = json.loads(capsys.readouterr().out.strip().splitlines()[-1])
    assert out["status"] == "pending" and "couldn't be replaced" in out["message"]
    assert xlsx().read_bytes() == before and X.pending_path(xlsx()).exists()


def test_disabled_and_unwritable(xconn, monkeypatch, tmp_path):
    conn, a, b = xconn
    from jobsearch import config
    prefs = config.load_preferences()
    prefs.excel_applications.enabled = False
    monkeypatch.setattr(config, "load_preferences", lambda *a, **k: prefs)
    assert X.update(conn) == "" and not xlsx().exists()
    prefs.excel_applications.enabled = True
    ro = tmp_path / "ro"
    ro.mkdir()
    os.chmod(ro, 0o500)
    try:
        msg = X.update(conn, ro / "sub" / "A.xlsx")
    finally:
        os.chmod(ro, 0o700)
    assert "macOS did not let" in msg and json.loads(db.get_meta(conn, "excel_applications"))["status"] == "error"
