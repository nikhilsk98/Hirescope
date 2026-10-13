"""Phase 6 (DEMO data): lock, notifications, source health, quick plan, Excel, backups/restore, digest, seasons,
launchd schedule rendering, offline handling, run summary."""
import gzip
import json
import os
import plistlib
import sqlite3
from datetime import date, datetime, timedelta, timezone
from types import SimpleNamespace

import httpx
import pytest
from openpyxl import load_workbook

from jobsearch.config import Company
from jobsearch.ops import alerts, backup, digest, excel, health, notify, schedule, seasons
from jobsearch.ops.lock import Locked, RunLock
from jobsearch.pipeline import plan
from jobsearch.storage import db
from jobsearch.tracker import statuses

from .test_storage import rec


# --- lock ------------------------------------------------------------------------------------------
def test_lock_blocks_second_run_and_releases(tmp_path):
    p = tmp_path / "run.lock"
    with RunLock(p):
        assert json.loads(p.read_text())["pid"] == os.getpid()
        with pytest.raises(Locked):
            RunLock(p).acquire()
    assert not p.exists()


def test_stale_locks_are_replaced(tmp_path):
    p = tmp_path / "run.lock"
    p.write_text(json.dumps({"pid": 999_999_9, "started": datetime.now(timezone.utc).isoformat()}))   # dead pid
    with RunLock(p):
        assert json.loads(p.read_text())["pid"] == os.getpid()
    old = (datetime.now(timezone.utc) - timedelta(hours=5)).isoformat()
    p.write_text(json.dumps({"pid": os.getppid(), "started": old}))                             # alive but hung
    with RunLock(p, stale_hours=3):
        pass
    p.write_text("garbage")
    with RunLock(p):
        pass


# --- notifications ---------------------------------------------------------------------------------
def test_notification_command_escapes_job_text(monkeypatch):
    cmd = notify.command('3 new "Strong" matches', 'DemoCo — SWE \\ "New Grad"\nline2', "sub")
    assert cmd[:2] == ["osascript", "-e"]
    assert '\\"Strong\\"' in cmd[2] and '\\\\' in cmd[2] and "\n" not in cmd[2]
    calls = []
    monkeypatch.setattr(notify.sys, "platform", "darwin")
    monkeypatch.delenv("HIRESCOPE_NO_NOTIFY", raising=False)
    monkeypatch.setattr(notify, "ENABLED", True)
    assert notify.notify("t", "m", runner=lambda c, **k: calls.append(c)) and calls
    monkeypatch.setattr(notify, "ENABLED", False)
    assert not notify.notify("t", "m", runner=lambda c, **k: calls.append(c)) and len(calls) == 1


def test_alerts_send_and_summary(monkeypatch):
    sent = []
    monkeypatch.setattr(notify, "notify", lambda *a, **k: sent.append(a))
    s = {"new_strong_rows": [{"company": "DemoCo", "title": "SWE (DEMO)", "score": 88.0}] * 3,
         "early_bird": [{"company": "DemoCo", "title": "SWE (DEMO)", "score": 88.0}],
         "reopened": ["DemoCo: SWE (DEMO)"], "health_alerts": [{"message": "DemoCo returned 0 jobs"}],
         "sources_ok": 3, "sources_checked": 4, "by_track": {"full_time": {"visible": 3, "strong": 3, "stretch": 0}}}
    titles = alerts.send(s)
    assert len(titles) == 3 and "3 new Strong matches" in titles[0] and sent[0][2].startswith("1 early bird")
    txt = alerts.summary_text(s)
    assert "3 of 4 sources OK" in txt and "HEALTH: DemoCo returned 0 jobs" in txt and "REOPENED" in txt


# --- source health ---------------------------------------------------------------------------------
def test_health_thresholds():
    o = lambda c, n, ok=True, unchanged=False: SimpleNamespace(company=c, source="greenhouse", ok=ok, jobs=n, unchanged=unchanged)  # noqa: E731
    hist = {"Zero": [20, 22, 18], "Drop": [100, 110, 90], "Fine": [100, 95], "Small": [6, 5, 7], "New": [50],
            "Tiny": [1, 2], "Down": [100, 100]}
    alerts_ = health.check([o("Zero", 0), o("Drop", 25), o("Fine", 40), o("Small", 1), o("New", 0), o("Tiny", 0),
                            o("Down", 0, ok=False), o("Drop", 0, unchanged=True)], lambda c, s: hist[c])
    got = {(a["company"], a["kind"]) for a in alerts_}
    assert got == {("Zero", "zero"), ("Drop", "drop")}          # Fine = -60% (< 70%), Small/Tiny too small, New no history


# --- quick plan ------------------------------------------------------------------------------------
def test_quick_plan_is_my_companies_plus_tier1(conn, cfg):
    cfg.companies += [Company(name="Tier2Co", source="greenhouse", board="t2", tier=2),
                      Company(name="Tier3Co", source="greenhouse", board="t3", tier=3)]
    from jobsearch.watchlist import WatchItem
    planned, _ = plan(conn, cfg, [WatchItem("MineCo")], quick=True)
    names = {(p.company.name if p.company else p.watch.company) for p in planned}
    assert names == {"DemoCo", "MineCo"}


# --- excel -----------------------------------------------------------------------------------------
def test_excel_export(conn, tmp_path):
    pk, _ = db.upsert_job(conn, rec(title="=HYPERLINK(\"x\") (DEMO)"))
    db.save_scores(conn, pk, filtered=0, score=81, qual_score=70, label="Strong")
    conn.execute("UPDATE jobs SET track='full_time' WHERE id=?", (pk,))
    statuses.merge(conn, {"exported_at": "2026-10-12T00:00:00Z", "format": 3, "statuses": {},
                          "applications": {"a1": {"company": "DemoCo", "title": "SWE", "applied_date": "2026-10-01",
                                                  "updated_at": "2026-10-01T00:00:00Z"}}}, None)
    out = excel.export(conn, tmp_path / "x.xlsx", 100)
    wb = load_workbook(out)
    assert wb.sheetnames == ["Co-op & Internships", "Full-time", "Applications", "Coverage"]
    ft = wb["Full-time"]
    assert ft["A2"].value == "DemoCo" and ft["B2"].value.startswith("'=")      # never a live formula
    assert ft["O2"].hyperlink.target == "https://example.com/1"
    assert wb["Applications"]["A2"].value == "DemoCo"


# --- backups ---------------------------------------------------------------------------------------
def _file_db(path, title):
    c = db.connect(path)
    db.upsert_job(c, rec(title=title))
    c.execute("INSERT INTO embeddings (text_hash, model, vector) VALUES ('h', 'm', x'00')")
    c.commit()
    return c


def test_backup_restore_and_prune(tmp_path, monkeypatch):
    monkeypatch.setattr(backup.paths, "DATA", tmp_path / "data")
    live = tmp_path / "jobs.db"
    c = _file_db(live, "Before (DEMO)")
    c.close()
    imports = tmp_path / "imports"
    imports.mkdir()
    (imports / "statuses-1.json").write_text("{}")
    root = tmp_path / "backups"
    b = backup.backup(live, root, "2026-10-01", keep=14, extra_dirs=(imports,))
    assert (b / "jobs.db.gz").exists() and (b / "imports" / "statuses-1.json").exists()
    assert backup.backup(live, root, "2026-10-01") is None                     # once a day
    snap = sqlite3.connect(":memory:")
    snap.deserialize(gzip.decompress((b / "jobs.db.gz").read_bytes()))
    assert snap.execute("SELECT COUNT(*) FROM embeddings").fetchone()[0] == 0   # cache dropped
    c = db.connect(live)
    c.execute("UPDATE jobs SET title='After (DEMO)'")
    c.commit()
    c.close()
    kept = backup.restore(root, "2026-10-01", live)
    assert kept.name.startswith("jobs.db.before-restore-")
    assert sqlite3.connect(live).execute("SELECT title FROM jobs").fetchone()[0] == "Before (DEMO)"
    with pytest.raises(FileNotFoundError):
        backup.restore(root, "1999-01-01", live)
    for d in range(2, 20):
        backup.backup(live, root, f"2026-10-{d:02d}", keep=14)
    names = [p.name for p in backup.list_backups(root)]
    assert len(names) == 14 and names[0] == "2026-10-06" and names[-1] == "2026-10-19"


# --- digest ----------------------------------------------------------------------------------------
def test_digest_escapes_and_summarizes(conn, tmp_path):
    pk, _ = db.upsert_job(conn, rec(title="<script>alert(1)</script> Engineer (DEMO)"))
    db.save_scores(conn, pk, filtered=0, score=90, label="Strong")
    conn.execute("UPDATE jobs SET track='full_time', first_seen=? WHERE id=?", (datetime.now(timezone.utc).isoformat(), pk))
    today = datetime.now(timezone.utc).date()
    statuses.merge(conn, {"exported_at": datetime.now(timezone.utc).isoformat(), "format": 3, "statuses": {},
                          "applications": {"a1": {"company": "DemoCo", "title": "SWE", "applied_date": (today - timedelta(days=8)).isoformat(),
                                                  "updated_at": "2026-10-01T00:00:00Z", "track": "full_time"}}}, None)
    html = digest.build(conn, today, [{"company": "DemoCo", "track": "coop_intern", "status": "expected in 9 days",
                                       "confidence": "one year of history", "remind": True}])
    assert "<script>" not in html and "&lt;script&gt;" in html
    assert "New Strong matches — Full-time (1)" in html and "Follow-ups due (1)" in html and "Hiring seasons" in html
    assert digest.due(date(2026, 10, 11), "Sunday", tmp_path) and not digest.due(date(2026, 10, 12), "Sunday", tmp_path)
    p = digest.write(conn, tmp_path, date(2026, 10, 11))
    assert p.name == "2026-10-11.html" and not digest.due(date(2026, 10, 11), "Sunday", tmp_path)


# --- seasons ---------------------------------------------------------------------------------------
def test_hiring_seasons(conn):
    def add(jid, title, posted, track, status="closed"):
        pk, _ = db.upsert_job(conn, rec(job_id=jid, title=title, normalized_title=title.lower(), posted_date=posted))
        conn.execute("UPDATE jobs SET track=?, job_status=? WHERE id=?", (track, status, pk))
    add("1", "SWE Co-op (DEMO)", "2025-09-03T00:00:00+00:00", "coop_intern")
    add("2", "SWE Co-op Spring (DEMO)", "2025-09-20T00:00:00+00:00", "coop_intern")
    add("3", "SWE Co-op (DEMO)", "2024-09-10T00:00:00+00:00", "coop_intern")
    add("4", "Software Engineer, New Grad (DEMO)", "2025-08-01T00:00:00+00:00", "full_time")
    add("5", "Software Engineer (DEMO)", "2025-05-01T00:00:00+00:00", "full_time")       # not new grad: ignored
    out = {(s["company"], s["track"]): s for s in seasons.track_seasons(conn, date(2026, 8, 20))}
    co = out[("DemoCo", "coop_intern")]
    assert co["years"] == {"2024": "2024-09-10", "2025": "2025-09-03"} and co["predicted"] == "2026-09"
    assert co["remind"] and co["status"] == "expected in 14 days" and co["confidence"] == "repeating pattern"
    ng = out[("DemoCo", "new_grad")]
    assert ng["predicted"] == "2026-08" and not ng["remind"] and ng["confidence"] == "one year of history"
    assert len(out) == 2


# --- launchd schedule ------------------------------------------------------------------------------
def test_schedule_render_and_install(tmp_path):
    full = plistlib.loads(schedule.render_full("18:00,08:00", repo=tmp_path / "repo & co", log_dir=tmp_path).encode())
    assert full["Label"] == "com.nikhil.jobsearch"
    assert full["StartCalendarInterval"] == [{"Hour": 8, "Minute": 0}, {"Hour": 18, "Minute": 0}]
    assert full["ProgramArguments"][1].endswith("repo & co/scripts/run.sh")
    quick = plistlib.loads(schedule.render_quick(2.5, repo=tmp_path, log_dir=tmp_path).encode())
    assert quick["StartInterval"] == 9000 and quick["ProgramArguments"][-1] == "--quick"
    for bad in ("25:00", "8", "", "08:60"):
        with pytest.raises(ValueError):
            schedule.parse_times(bad)
    with pytest.raises(ValueError):
        schedule.render_quick(0.1)
    calls = []
    run = lambda cmd, **k: (calls.append(cmd), SimpleNamespace(returncode=0))[1]  # noqa: E731
    written = schedule.install(["07:30"], 3, agents=tmp_path / "agents", run=run)
    assert [p.name for p in written] == ["com.nikhil.jobsearch.plist", "com.nikhil.jobsearch.quick.plist"]
    assert [c[1] for c in calls] == ["bootout", "bootout", "bootstrap", "bootstrap"]
    schedule.install(["07:30"], 0, agents=tmp_path / "agents", run=run)          # quick checks off
    assert not (tmp_path / "agents" / "com.nikhil.jobsearch.quick.plist").exists()
    schedule.uninstall(agents=tmp_path / "agents", run=run)
    assert not list((tmp_path / "agents").iterdir())


# --- offline / lock in main ------------------------------------------------------------------------
def test_main_offline_and_locked(monkeypatch, tmp_path):
    from jobsearch import run as R
    monkeypatch.setattr(R.paths, "LOG_DIR", tmp_path / "logs")
    monkeypatch.setattr(R.paths, "DATA", tmp_path)
    monkeypatch.setattr(R.alerts, "online", lambda: False)
    assert R.main([]) == R.EX_TEMPFAIL
    monkeypatch.setattr(R.alerts, "online", lambda: True)
    with RunLock(tmp_path / "run.lock"):
        assert R.main(["--quick"]) == 0          # a quick check yields to a running full run
        assert R.main([]) == 3


def test_online_probe():
    assert alerts.online(timeout=0.2, hosts=(("127.0.0.1", 9),)) is False


# --- alerts: early birds include Stretch, season reminders once, exit codes -------------------------
def test_early_birds_include_stretch(conn):
    now = datetime.now(timezone.utc)
    def add(jid, label, posted_h, score):
        pk, _ = db.upsert_job(conn, rec(job_id=jid, title=f"SWE {jid} (DEMO)", normalized_title=f"swe {jid}",
                                        posted_date=(now - timedelta(hours=posted_h)).isoformat()))
        conn.execute("UPDATE jobs SET label=?, score=?, filtered=0 WHERE id=?", (label, score, pk))
    add("1", "Strong", 10, 90)
    add("2", "Stretch", 5, 60)       # fresh Stretch: an early bird, not a "new Strong match"
    add("3", "Stretch", 100, 65)     # too old for an early bird
    add("4", "Weak", 2, 30)
    strong, early = alerts.new_strong(conn, (now - timedelta(days=1)).isoformat(), 48)
    assert [r["title"] for r in strong] == ["SWE 1 (DEMO)"]
    assert [r["title"] for r in early] == ["SWE 1 (DEMO)", "SWE 2 (DEMO)"]
    sent = []
    notify_orig, notify.notify = notify.notify, lambda *a, **k: sent.append(a)
    try:
        titles = alerts.send({"new_strong_rows": [], "early_bird": [early[1]]})
    finally:
        notify.notify = notify_orig
    assert titles == ["HireScope: 1 early bird"] and "Strong/Stretch" in sent[0][2]


def test_season_reminders_notified_once(conn, monkeypatch):
    sent = []
    monkeypatch.setattr(notify, "notify", lambda *a, **k: sent.append(a))
    sea = [{"company": "DemoCo", "track": "coop_intern", "predicted": "2026-11", "remind": True, "status": "expected in 9 days"},
           {"company": "OtherCo", "track": "new_grad", "predicted": "2027-02", "remind": False, "status": "expected around Feb 2027"}]
    first = alerts.season_reminders(conn, sea)
    assert [x["company"] for x in first] == ["DemoCo"]
    assert alerts.season_reminders(conn, sea) == []                  # remembered: no second notification
    assert alerts.send({"season_reminders": first}) == ["HireScope: 1 hiring season opening soon"]
    assert "DemoCo (co-op/intern, in 9 days)" in sent[0][1]
    assert "SEASON: DemoCo" in alerts.summary_text({"season_reminders": first})


def test_exit_code_ignores_manual_companies():
    from jobsearch.run import exit_code
    assert exit_code({"failures": [], "manual": [{"company": "Apple"}], "partial": [{"company": "RTX"}]}) == 0
    assert exit_code({"failures": [{"company": "X", "source": "greenhouse", "error": "404"}], "manual": []}) == 1


def test_quick_check_skips_the_aggregator(conn, cfg, monkeypatch):
    from jobsearch.adapters.aggregator import AggregatorAdapter
    from jobsearch.config import Aggregator
    monkeypatch.setenv("JSEARCH_API_KEY", "demo-key")
    monkeypatch.setattr(AggregatorAdapter, "enabled", True)
    prefs = cfg.prefs.model_copy(update={"aggregator": Aggregator(enabled=True)})
    cfg2 = cfg.model_copy(update={"prefs": prefs})
    full = {p.company.source for p in plan(conn, cfg2, [], force_all=True)[0]}
    quick = {p.company.source for p in plan(conn, cfg2, [], quick=True)[0]}
    assert "aggregator" in full and "aggregator" not in quick           # free-tier quotas: full runs only


# --- end to end: python -m jobsearch.run with mocked sources (DEMO) --------------------------------
def test_main_end_to_end_with_mocked_sources(monkeypatch, tmp_path, cfg, greenhouse_payload):
    """The whole run: plan -> fetch (mocked Greenhouse) -> extract -> score -> SQLite -> dashboard data.js/sync.js ->
    run summary, all inside tmp_path (the real data/ and dashboard/ are never touched)."""
    from jobsearch import paths
    from jobsearch import run as R
    from .conftest import FIXTURES, FakeEmbedder, mock_client
    from .test_pipeline import gh_handler
    for name, rel in (("DATA", ""), ("LOG_DIR", "logs"), ("JD_DIR", "jd"), ("APPS_DIR", "applications"),
                      ("IMPORTS_DIR", "imports"), ("PACKAGES_DIR", "packages")):
        monkeypatch.setattr(paths, name, tmp_path / rel if rel else tmp_path)
    monkeypatch.setattr(paths, "DB", tmp_path / "jobs.db")
    monkeypatch.setattr(paths, "ROOT", tmp_path)
    monkeypatch.setattr(paths, "DASHBOARD", tmp_path / "dashboard")
    monkeypatch.setattr(paths, "SYNC_JS", tmp_path / "dashboard" / "sync.js")
    (tmp_path / "inputs").mkdir()
    (tmp_path / "inputs" / "my_companies.csv").write_text("company,careers_url,notes\nWatchCo,https://jobs.lever.co/watchco,DEMO\n")
    prefs = cfg.prefs.model_copy(update={"statuses_import_dir": str(tmp_path / "downloads"), "my_companies_csv": "inputs/my_companies.csv"})
    monkeypatch.setattr(R, "load_config", lambda *a, **k: cfg.model_copy(update={"prefs": prefs}))
    lever = [{"id": "22222222-2222-2222-2222-222222222222", "text": "Software Engineer, New Grad (DEMO)",
              "categories": {"commitment": "Full-time", "location": "Boston, MA", "allLocations": ["Boston, MA"]},
              "createdAt": 1791676800000, "description": "<p>DEMO</p>", "lists": [], "additional": "",
              "hostedUrl": "https://jobs.lever.co/watchco/22222222-2222-2222-2222-222222222222"}]
    gh = gh_handler({"democo": greenhouse_payload})

    def handler(request):
        return httpx.Response(200, json=lever) if "api.lever.co" in str(request.url) else gh(request)
    monkeypatch.setattr(R, "HttpClient", lambda **k: mock_client(handler))
    monkeypatch.setattr(R, "Embedder", FakeEmbedder)
    monkeypatch.setattr(R.alerts, "online", lambda: True)
    args = ["--no-email", "--no-notify", "--no-backup", "--no-link-check", "--no-import",
            "--resume", str(FIXTURES / "demo_resume.md"), "--prefs", str(FIXTURES / "demo_preferences.yaml")]
    assert R.main(args) == 0
    data = (tmp_path / "dashboard" / "data.js").read_text()
    assert "window.META" in data and "Software Engineer, New Grad 2027 (DEMO)" in data and "WatchCo" in data
    assert (tmp_path / "dashboard" / "sync.js").exists()
    summary = (tmp_path / "logs" / "last_run_summary.txt").read_text()
    assert "2 of 2 sources OK" in summary and "0 failed" in summary
    import sqlite3
    c = sqlite3.connect(tmp_path / "jobs.db")
    assert c.execute("SELECT COUNT(*) FROM jobs").fetchone()[0] == 4
    assert c.execute("SELECT source FROM company_sources WHERE company='WatchCo'").fetchone()[0] == "lever"
    # a dry run afterwards writes nothing
    before = c.execute("SELECT COUNT(*) FROM runs").fetchone()[0]
    c.close()
    assert R.main([*args, "--dry-run"]) == 0
    c = sqlite3.connect(tmp_path / "jobs.db")
    assert c.execute("SELECT COUNT(*) FROM runs").fetchone()[0] == before
